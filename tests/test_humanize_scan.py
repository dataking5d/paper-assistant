import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import docx

from core.docx_parser import document_snapshot, text_hash
from core.humanize_scanner import scan_docx, scan_text

ROOT = Path(__file__).resolve().parents[1]


class TestHumanizeScan(unittest.TestCase):
    def test_scan_text_reports_candidates_without_replacement(self):
        result = scan_text("这不是一个工具，而是一个平台。研究表明，该方案有效。")
        ids = {f["rule_id"] for f in result["findings"]}
        self.assertIn("zh.pattern.01.not_x_but_y", ids)
        self.assertIn("zh.pattern.17.vague_authority", ids)
        for f in result["findings"]:
            self.assertFalse(f["auto_fixable"])
            self.assertTrue(f["requires_agent"])
            self.assertNotIn("replacement", f)

    def test_protected_numbers_escalate(self):
        result = scan_text("不是降低 10%，而是提高 5%。")
        hits = [f for f in result["findings"] if f["rule_id"] == "zh.pattern.01.not_x_but_y"]
        self.assertTrue(hits)
        self.assertTrue(hits[0]["protected"])
        self.assertEqual(hits[0]["severity"], "error")

    def test_docx_scan_is_read_only_and_matches_parse_anchors(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "paper.docx"
            d = docx.Document()
            d.add_heading("一、测试", level=1)
            d.add_paragraph("这不是一个工具，而是一个平台。研究表明，该方案有效。")
            d.save(path)
            before = hashlib.sha256(path.read_bytes()).hexdigest()
            snapshot = document_snapshot(path)
            result = scan_docx(path)
            after = hashlib.sha256(path.read_bytes()).hexdigest()
            self.assertEqual(before, after)
            self.assertEqual(result["document_sha256"], before)
            self.assertEqual(len(result["paragraphs"]), 1)
            para = result["paragraphs"][0]
            self.assertEqual(para["paragraph_index"], 1)
            self.assertEqual(para["section"], "一、测试")
            self.assertEqual(para["text_sha256"], text_hash("这不是一个工具，而是一个平台。研究表明，该方案有效。"))
            self.assertEqual(para["text_sha256"], snapshot["paragraphs"][1]["text_sha256"])

    def test_cli_returns_json_and_does_not_modify_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "paper.docx"
            d = docx.Document()
            d.add_paragraph("这不是一个工具，而是一个平台。")
            d.save(path)
            before = hashlib.sha256(path.read_bytes()).hexdigest()
            proc = subprocess.run(
                [sys.executable, "main.py", "humanize-scan", str(path), "--json"],
                cwd=ROOT, capture_output=True, text=True, encoding="utf-8",
            )
            self.assertEqual(proc.returncode, 0, proc.stderr)
            payload = json.loads(proc.stdout)
            self.assertEqual(payload["schema_version"], 1)
            self.assertGreaterEqual(payload["summary"]["findings"], 1)
            self.assertEqual(before, hashlib.sha256(path.read_bytes()).hexdigest())


if __name__ == "__main__":
    unittest.main()
