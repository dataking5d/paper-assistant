import copy
import hashlib
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt

from core.docx_parser import document_snapshot
from core.agent_workflow import apply_edit_plan, read_evidence


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / '原件.docx'
        doc = Document()
        doc.add_heading('1. Introduction', 1)
        doc.add_paragraph('Original findings [1].')
        doc.add_heading('1. Introduction', 1)
        doc.add_paragraph('Another paragraph.')
        doc.add_paragraph('参考文献')
        doc.add_paragraph('[1] Able. Original study[J]. Journal, 2020.')
        doc.add_paragraph('致 谢')
        doc.add_paragraph('KEEP TAIL')
        doc.save(self.source)
        self.evidence = self.root / 'source.txt'
        self.evidence.write_text('The study found better clarity.', encoding='utf-8')
        self.snapshot = document_snapshot(str(self.source))

    def plan(self, operations=None):
        return {'schema_version': 1, 'source_sha256': self.snapshot['source_sha256'],
                'operations': operations or [self.op('replace', 1, text='Clearer findings [1].')],
                'references': {}}

    def op(self, kind, index, **kwargs):
        return {'op': kind, 'paragraph_index': index,
                'text_sha256': self.snapshot['paragraphs'][index]['text_sha256'], **kwargs}

    def run_plan(self, plan, name='attempt', validate=False):
        return apply_edit_plan(str(self.source), plan, str(self.root / name),
                               base_dir=self.root, validate_only=validate)

    def test_plain_polish_preserves_input_and_tail(self):
        before = self.source.read_bytes()
        result = self.run_plan(self.plan())
        self.assertEqual(result['status'], 'ready_for_review')
        self.assertEqual(self.source.read_bytes(), before)
        output = Document(result['artifacts']['document'])
        self.assertEqual(output.paragraphs[1].text, 'Clearer findings [1].')
        self.assertEqual(output.paragraphs[-1].text, 'KEEP TAIL')
        self.assertFalse(result['checks']['user_accepted'])
        self.assertTrue(Path(result['artifacts']['report']).exists())

    def test_validate_only_no_output(self):
        result = self.run_plan(self.plan(), validate=True)
        self.assertEqual(result['status'], 'validated')
        self.assertFalse((self.root / 'attempt').exists())

    def test_stale_snapshot_unknown_target_and_empty_plan(self):
        for alteration in ['stale', 'target', 'empty']:
            plan = self.plan()
            if alteration == 'stale':
                plan['source_sha256'] = '0' * 64
            elif alteration == 'target':
                plan['operations'][0]['paragraph_index'] = 100
            else:
                plan['operations'] = []
            with self.subTest(alteration=alteration), self.assertRaises(ValueError):
                self.run_plan(plan)
            self.assertFalse((self.root / 'attempt').exists())

    def test_protect_citations_and_numbers(self):
        plan = self.plan()
        plan['operations'][0]['text'] = 'Fabricated 77 findings [2].'
        with self.assertRaises(ValueError):
            self.run_plan(plan)

    def citation_plan(self):
        plan = self.plan([self.op('insert_citation', 3, text='Supported finding.', ref_id='r1')])
        evidence = read_evidence(str(self.evidence))
        plan['references']['r1'] = {'entry': 'Baker. Clarity study[J]. Journal, 2021.',
            'source': 'source.txt', 'source_sha256': evidence['source_sha256'],
            'page': 1, 'quote': 'better clarity', 'support_note': 'The source describes clarity.'}
        return plan

    def test_citation_and_combined_workflow(self):
        plan = self.citation_plan()
        plan['operations'].insert(0, self.op('replace', 1, text='Clearer findings [1].'))
        plan['operations'].append(self.op('format', 3, formatting={'font_cn': '宋体', 'font_en': 'Arial', 'size_pt': 12, 'line_spacing': 1.5}))
        with patch('socket.socket', side_effect=AssertionError('No network allowed')):
            result = self.run_plan(plan)
        text = [p.text for p in Document(result['artifacts']['document']).paragraphs]
        self.assertIn('Supported finding.[2]', text)
        self.assertIn('[1] Able. Original study[J]. Journal, 2020.', text)
        self.assertEqual(result['applied'], 3)

    def test_unknown_reference_and_false_quote(self):
        for mutation in ['ref', 'quote', 'hash']:
            plan = self.citation_plan()
            if mutation == 'ref':
                plan['operations'][0]['ref_id'] = 'absent'
            elif mutation == 'quote':
                plan['references']['r1']['quote'] = 'The opposite result.'
            else:
                plan['references']['r1']['source_sha256'] = '0' * 64
            with self.subTest(mutation=mutation), self.assertRaises(ValueError):
                self.run_plan(plan)

    def test_existing_attempt_refused(self):
        self.run_plan(self.plan())
        with self.assertRaises(ValueError):
            self.run_plan(self.plan())

    def test_complex_paragraph_refused(self):
        doc = Document(self.source)
        doc.paragraphs[1]._p.append(OxmlElement('w:fldSimple'))
        doc.save(self.source)
        self.snapshot = document_snapshot(str(self.source))
        with self.assertRaises(ValueError):
            self.run_plan(self.plan())

    def test_mixed_run_format_refused_in_polishing(self):
        doc = Document(self.source)
        doc.paragraphs[3].add_run('Important').bold = True
        doc.save(self.source)
        self.snapshot = document_snapshot(str(self.source))
        with self.assertRaises(ValueError):
            self.run_plan(self.plan([self.op('replace', 3, text='Reworded Important')]))

    def test_cli_json_without_api_or_optional_packages(self):
        plan_path = self.root / 'edits.json'
        plan_path.write_text(json.dumps(self.plan()), encoding='utf-8')
        env = {k: v for k, v in os.environ.items() if k not in ('DEEPSEEK_API_KEY', 'ANTHROPIC_AUTH_TOKEN')}
        env['PYTHONIOENCODING'] = 'utf-8'
        result = subprocess.run([sys.executable, '-B', 'main.py', 'edit', '--paper', str(self.source),
            '--edits', str(plan_path), '--attempt', str(self.root / 'cli')], capture_output=True, text=True,
            encoding='utf-8', env=env)
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        self.assertEqual(json.loads(result.stdout)['status'], 'ready_for_review')

    def test_polish_preserves_untouched_parts(self):
        import zipfile
        doc = Document(self.source)
        doc.sections[0].header.paragraphs[0].text = 'UNTOUCHED HEADER'
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Untouched table'
        doc.save(self.source)
        self.snapshot = document_snapshot(str(self.source))
        result = self.run_plan(self.plan())
        with zipfile.ZipFile(self.source) as old, zipfile.ZipFile(result['artifacts']['document']) as new:
            self.assertEqual(old.namelist(), new.namelist())
            for name in old.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(old.read(name), new.read(name), name)
        self.assertEqual(Document(result['artifacts']['document']).tables[0].cell(0, 0).text, 'Untouched table')

    def test_failed_batch_does_not_publish_partial_result(self):
        plan = self.plan()
        plan['operations'].append(self.op('replace', 3, text='Unexpected 999'))
        with self.assertRaises(ValueError):
            self.run_plan(plan)
        self.assertFalse((self.root / 'attempt').exists())

    def test_reference_format_only(self):
        plan = self.plan()
        plan['operations'] = []
        plan['reference_format'] = {'size_pt': 10.5, 'first_line_indent_pt': -21}
        result = self.run_plan(plan)
        para = next(p for p in Document(result['artifacts']['document']).paragraphs if p.text.startswith('[1]'))
        self.assertEqual(para.paragraph_format.left_indent.pt, 21)
        self.assertEqual(para.runs[0].font.size.pt, 10.5)

    def test_pdf_physical_pages_and_quote_validation(self):
        # A two-page PDF built with stdlib; page 2 deliberately contains no text.
        content = b'BT /F1 12 Tf 72 700 Td (The study found better clarity.) Tj ET'
        objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
            b'<< /Type /Pages /Kids [3 0 R 5 0 R] /Count 2 >>',
            b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 7 0 R >> >> /Contents 4 0 R >>',
            b'<< /Length ' + str(len(content)).encode() + b' >>\nstream\n' + content + b'\nendstream',
            b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << >> /Contents 6 0 R >>',
            b'<< /Length 0 >>\nstream\n\nendstream',
            b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
        raw = bytearray(b'%PDF-1.4\n')
        offsets = [0]
        for i, obj in enumerate(objects, 1):
            offsets.append(len(raw))
            raw += f'{i} 0 obj\n'.encode() + obj + b'\nendobj\n'
        xref = len(raw)
        raw += f'xref\n0 {len(offsets)}\n0000000000 65535 f \n'.encode()
        for offset in offsets[1:]:
            raw += f'{offset:010d} 00000 n \n'.encode()
        raw += f'trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode()
        path = self.root / 'source.pdf'
        path.write_bytes(raw)
        evidence = read_evidence(str(path))
        self.assertEqual(len(evidence['pages']), 2)
        self.assertEqual(evidence['pages'][1]['text'], '')
        plan = self.citation_plan()
        plan['references']['r1'].update(source='source.pdf', source_sha256=evidence['source_sha256'])
        self.assertEqual(self.run_plan(plan)['status'], 'ready_for_review')
        plan['references']['r1']['page'] = 2
        with self.assertRaises(ValueError):
            self.run_plan(plan, name='wrong-page')

    def test_conflicting_and_bad_format_refused(self):
        for fmt in [{'size_pt': float('nan')}, {'fake_setting': 10}]:
            with self.subTest(fmt=fmt), self.assertRaises(ValueError):
                self.run_plan(self.plan([self.op('format', 3, formatting=fmt)]))

    def test_schema_typo_is_not_silently_ignored(self):
        plan = self.plan()
        plan['operations'][0]['paragraph_id'] = 1
        with self.assertRaises(ValueError):
            self.run_plan(plan)

    def test_cli_failure_is_structured_and_no_artifact(self):
        plan = self.plan()
        plan['source_sha256'] = 'outdated'
        path = self.root / 'bad.json'
        path.write_text(json.dumps(plan), encoding='utf-8')
        result = subprocess.run([sys.executable, '-B', 'main.py', 'edit', '--paper', str(self.source),
            '--edits', str(path), '--attempt', str(self.root / 'rejected')], capture_output=True,
            text=True, encoding='utf-8')
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(json.loads(result.stdout)['status'], 'needs_input')
        self.assertFalse((self.root / 'rejected').exists())

    def test_noop_format_refused(self):
        doc = Document(self.source)
        doc.paragraphs[3].runs[0].font.size = Pt(12)
        doc.save(self.source)
        self.snapshot = document_snapshot(str(self.source))
        with self.assertRaises(ValueError):
            self.run_plan(self.plan([self.op('format', 3, formatting={'size_pt': 12})]))

    def test_two_attempts_cannot_publish_to_same_directory(self):
        from concurrent.futures import ThreadPoolExecutor
        def run():
            try:
                return self.run_plan(self.plan())['status']
            except (ValueError, FileExistsError):
                return 'rejected'
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: run(), range(2)))
        self.assertEqual(sorted(results), ['ready_for_review', 'rejected'])
        report = json.loads((self.root / 'attempt/report.json').read_text(encoding='utf-8'))
        self.assertEqual(report['applied'], 1)


if __name__ == '__main__':
    unittest.main()
