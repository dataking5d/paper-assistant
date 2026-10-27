import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Pt

from core.docx_parser import DocxParser
from core.docx_editor import DocxEditor, CitationEdit
from core.reference_kb import ReferenceKB, ReferenceEntry
from core.template_analyzer import _extract_paragraph_formatting


class RegressionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def paper(self, tail='致 谢'):
        doc = Document()
        for text in ['1. Introduction', 'Intro', '1.1 Alpha', 'Alpha [1].',
                     '1.2 Beta', 'Beta [2].', '参考文献',
                     '[1] Able. Alpha study[J]. Journal, 2020.',
                     '[2] Baker. Beta study[J]. Journal, 2021.', tail, 'KEEP TAIL']:
            doc.add_paragraph(text)
        path = self.root / '原论文.docx'
        doc.save(path)
        return path

    def test_numbered_hierarchy(self):
        paper = DocxParser(str(self.paper())).parse()
        self.assertEqual(len(paper.chapters), 1)
        self.assertEqual([s.heading for s in paper.chapters[0].sections], ['1.1 Alpha', '1.2 Beta'])

    def test_citation_range(self):
        self.assertEqual(DocxParser._extract_citations('Text [1, 3-5].'), ['1', '3', '4', '5'])

    def test_tail_boundaries(self):
        for title in ['致 谢', '附 录', 'Acknowledgments', 'ACKNOWLEDGMENTS']:
            with self.subTest(title=title):
                doc = Document(self.paper(title))
                refs, end, _ = DocxEditor._parse_old_refs(doc, 6)
                self.assertEqual(len(refs), 2)
                self.assertEqual(end, 9)

    def test_template_indent_units(self):
        doc = Document()
        p = doc.add_paragraph('reference')
        p.runs[0].font.size = Pt(12)
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.first_line_indent = Pt(-24)
        self.assertEqual(_extract_paragraph_formatting(p)['indent_value'], '2char')

    def test_kb_roundtrip_and_repeat_load(self):
        kb = ReferenceKB(str(self.root / 'kb'))
        kb.entries = [ReferenceEntry(title='Study', full_text='Actual source evidence.')]
        kb._save_index()
        loaded = ReferenceKB(str(self.root / 'kb'))
        loaded.load_index()
        self.assertEqual(loaded.entries[0].full_text, 'Actual source evidence.')
        loaded.load_index()
        self.assertEqual(len(loaded.entries), 1)

    def test_existing_refs_and_placement(self):
        source = self.paper()
        output = self.root / 'out.docx'
        edits = [CitationEdit('1.1 Alpha', 'New evidence', 'Gamma study', 'Clark', '2022',
                              'Clark. Gamma study[J]. Journal, 2022.'),
                 CitationEdit('1.2 Beta', 'Other evidence', 'Alpha study', 'Able', '2020',
                              'Able. Alpha study[J]. Journal, 2020.')]
        DocxEditor().apply_edits(str(source), edits, str(output))
        text = [p.text for p in Document(output).paragraphs]
        self.assertIn('Beta [3].', text)
        self.assertIn('[3] Baker. Beta study[J]. Journal, 2021.', text)
        self.assertLess(text.index('1.2 Beta'), text.index('Other evidence[1]'))
        self.assertIn('KEEP TAIL', text)

    def test_same_path_refused(self):
        source = self.paper()
        before = source.read_bytes()
        with self.assertRaises(ValueError):
            DocxEditor().apply_edits(str(source), [], str(source))
        self.assertEqual(source.read_bytes(), before)

    def test_multirun_and_table_citations_keep_identity(self):
        source = self.paper()
        doc = Document(source)
        p = doc.paragraphs[3]
        p.clear()
        p.add_run('Alpha [').bold = True
        p.add_run('1').font.superscript = True
        p.add_run(',2].')
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'Table [2].'
        doc.paragraphs[6]._p.addprevious(table._tbl)
        doc.save(source)
        output = self.root / 'multirun.docx'
        edit = CitationEdit('1. Introduction', 'First addition', 'New', '', '', 'New. Study[J]. Journal, 2022.')
        DocxEditor().apply_edits(str(source), [edit], str(output))
        result = Document(output)
        self.assertIn('Alpha [2,3].', [p.text for p in result.paragraphs])
        self.assertEqual(result.tables[0].cell(0, 0).text, 'Table [3].')
        alpha = next(p for p in result.paragraphs if p.text.startswith('Alpha ['))
        self.assertTrue(alpha.runs[0].bold)
        self.assertTrue(alpha.runs[1].font.superscript)

    def test_ambiguous_bibliography_refused(self):
        source = self.paper('Unrecognized tail heading')
        with self.assertRaises(ValueError):
            DocxEditor().apply_edits(str(source), [CitationEdit('1.1 Alpha', 'Evidence', 'New', '', '', 'New entry')], str(self.root / 'out.docx'))
        self.assertFalse((self.root / 'out.docx').exists())

    def test_duplicate_heading_refused(self):
        source = self.paper()
        doc = Document(source)
        doc.paragraphs[4].text = '1.1 Alpha'
        doc.save(source)
        with self.assertRaises(ValueError):
            DocxEditor().apply_edits(str(source), [CitationEdit('1.1 Alpha', 'Evidence', 'New', '', '', 'New entry')], str(self.root / 'out.docx'))

    def test_old_kb_evidence_gap(self):
        import json
        root = self.root / 'old-kb'
        root.mkdir()
        (root / 'kb_index.json').write_text(json.dumps([{'title': 'Metadata only'}]), encoding='utf-8')
        kb = ReferenceKB(str(root))
        kb.load_index()
        self.assertIn('evidence_gap', kb.entries[0].meta)


if __name__ == '__main__':
    unittest.main()
