import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document
from core.citation_generator import CitationGenerator
from core.reference_kb import ReferenceEntry
from capabilities.writing.handler import _load_paper_context
from capabilities.polishing.handler import _extract_section_text
from capabilities.citation.handler import handle


class LegacyGuards(unittest.TestCase):
    def test_unsupported_request_does_not_fallback_to_citation(self):
        # Manifest discovery is outside this test; isolate the resolver from optional YAML.
        with patch.dict('sys.modules', {'yaml': types.ModuleType('yaml')}):
            from router.router import Router
            router = Router.__new__(Router)
            router._manifests = {'citation': {}, 'writing': {}, 'polishing': {}}
            for tasks in [['review'], ['review', 'citation'], ['writing', 'polishing']]:
                with self.subTest(tasks=tasks), self.assertRaises(ValueError):
                    router.resolve_pipeline(types.SimpleNamespace(task_types=tasks))

    def test_generation_exception_is_not_citation_text(self):
        client = Mock()
        client.messages.create.side_effect = RuntimeError('API failed')
        with self.assertRaises(RuntimeError):
            CitationGenerator(client).generate('Context', ReferenceEntry(title='Study'))

    def test_invalid_response_is_not_raw_citation(self):
        with self.assertRaises(ValueError):
            CitationGenerator(None)._parse_generation_response('not a result', ReferenceEntry(title='Study'))

    def test_section_context_reaches_handlers(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'paper.docx'
            doc = Document()
            doc.add_paragraph('1. Introduction')
            doc.add_paragraph('1.1 Background')
            doc.add_paragraph('ACTUAL EVIDENCE')
            doc.save(path)
            self.assertIn('ACTUAL EVIDENCE', _load_paper_context(str(path)))
            self.assertIn('ACTUAL EVIDENCE', _extract_section_text(str(path), 'full-paper'))

    def test_old_artifact_is_not_reported_as_new(self):
        with tempfile.TemporaryDirectory() as tmp:
            old = Path(tmp) / 'old_enhanced_v2.docx'
            old.write_bytes(b'old')
            fake = types.ModuleType('core.agent')
            fake.PaperAgent = Mock()
            fake.PaperAgent.return_value.run.return_value = None
            with patch.dict('sys.modules', {'core.agent': fake}):
                result = handle({'paper_path': str(Path(tmp) / 'paper.docx')})
            self.assertEqual(result['result_type'], 'error')
            self.assertFalse(result.get('artifacts'))


if __name__ == '__main__':
    unittest.main()
