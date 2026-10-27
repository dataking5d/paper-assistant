"""Offline file operations for a host Agent. No model clients or task scheduler."""

import hashlib
import json
import math
import os
import re
import shutil
import tempfile
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from .docx_editor import CitationEdit, DocxEditor
from .docx_parser import (CITATION_RE, is_reference_heading,
                          is_post_reference_heading, plain_paragraph, text_hash)
from .style_config import get_builtin_style


def read_evidence(source: str) -> dict:
    """Export actual local text with physical PDF page numbers (not printed labels)."""
    path = Path(source).resolve()
    raw = path.read_bytes()
    if path.suffix.lower() == '.pdf':
        import pdfplumber
        with pdfplumber.open(BytesIO(raw)) as pdf:
            pages = [{'page': i, 'text': page.extract_text() or ''}
                     for i, page in enumerate(pdf.pages, 1)]
    elif path.suffix.lower() in {'.txt', '.md'}:
        pages = [{'page': 1, 'text': raw.decode('utf-8-sig')}]
    else:
        raise ValueError('Evidence supports PDF and UTF-8 txt/md files')
    return {'source': str(path), 'source_sha256': hashlib.sha256(raw).hexdigest(),
            'pages': pages, 'warnings': ['No OCR or semantic truth verification is performed.']}


def _text(value, label):
    if not isinstance(value, str) or not value.strip() or any(c in value for c in '\r\n\x00'):
        raise ValueError(f'{label} must be a nonempty single-line string')
    return value


def _format(para, fmt):
    allowed = {'font_cn', 'font_en', 'size_pt', 'line_spacing', 'space_before_pt',
               'space_after_pt', 'first_line_indent_pt', 'left_indent_pt'}
    if not isinstance(fmt, dict) or not fmt or set(fmt) - allowed:
        raise ValueError('Unsupported or empty paragraph formatting')
    for key, value in fmt.items():
        if key.startswith('font_'):
            _text(value, key)
        elif type(value) not in (int, float) or not math.isfinite(value):
            raise ValueError(f'{key} must be a finite number')
        elif (key in {'size_pt', 'line_spacing'} and not 0 < value <= 100
              or key not in {'size_pt', 'line_spacing'} and not -1000 <= value <= 1000):
            raise ValueError(f'{key} is outside the supported range')
    for run in para.runs:
        if 'size_pt' in fmt:
            run.font.size = Pt(fmt['size_pt'])
        if 'font_en' in fmt:
            run.font.name = fmt['font_en']
        if 'font_cn' in fmt:
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), fmt['font_cn'])
    for key in ('space_before_pt', 'space_after_pt', 'first_line_indent_pt', 'left_indent_pt'):
        if key in fmt:
            setattr(para.paragraph_format, key[:-3], Pt(fmt[key]))
    if 'line_spacing' in fmt:
        para.paragraph_format.line_spacing = fmt['line_spacing']


def apply_edit_plan(source: str, plan: dict, attempt: str, *, base_dir=None,
                    validate_only: bool = False) -> dict:
    """Validate the entire batch, edit in memory, then publish one new attempt directory."""
    source_path, target = Path(source).resolve(), Path(attempt).resolve()
    base_dir = Path(base_dir or Path.cwd())
    if target.exists():
        raise ValueError('Attempt directory already exists; choose a new version')
    if not isinstance(plan, dict) or type(plan.get('schema_version')) is not int or plan['schema_version'] != 1:
        raise ValueError('Expected edit schema_version 1')
    if set(plan) - {'schema_version', 'source_sha256', 'operations', 'references', 'reference_format'}:
        raise ValueError('Unknown edit-plan fields')
    raw = source_path.read_bytes()
    source_hash = hashlib.sha256(raw).hexdigest()
    if plan.get('source_sha256') != source_hash:
        raise ValueError('Source file changed; parse it again')
    doc = Document(BytesIO(raw))
    original_xml = str(doc.element.xml)
    paras = doc.paragraphs
    operations = plan.get('operations')
    if not isinstance(operations, list) or (not operations and not plan.get('reference_format')):
        raise ValueError('At least one operation is required')
    references = plan.get('references', {})
    if not isinstance(references, dict):
        raise ValueError('references must be an object keyed by ref_id')
    ref_start = next((i for i, p in enumerate(paras) if is_reference_heading(p.text)), len(paras))
    protected_start = next((i for i, p in enumerate(paras) if is_post_reference_heading(p.text)), len(paras))
    positions, changes, seen, evidence_cache = [], [], set(), {}
    prepared = []
    for op in operations:
        if not isinstance(op, dict) or set(op) - {'op', 'paragraph_index', 'text_sha256', 'text', 'ref_id', 'formatting'}:
            raise ValueError('Invalid operation fields')
        kind, index = op.get('op'), op.get('paragraph_index')
        if kind not in {'replace', 'insert_citation', 'format'}:
            raise ValueError('Unsupported operation')
        if type(index) is not int or not 0 <= index < len(paras):
            raise ValueError('Unknown paragraph index')
        para = paras[index]
        if op.get('text_sha256') != text_hash(para.text):
            raise ValueError(f'Paragraph {index} changed')
        if index >= min(ref_start, protected_start) or not plain_paragraph(para):
            raise ValueError(f'Paragraph {index} is protected or contains unsupported structures')
        if kind != 'format' and para.style and para.style.name.startswith('Heading'):
            raise ValueError('Text changes require a body paragraph, not a heading')
        key = (kind, index)
        if key in seen and kind != 'insert_citation':
            raise ValueError('Conflicting duplicate operations')
        seen.add(key)
        if kind == 'replace':
            text = _text(op.get('text'), 'replacement text')
            if text == para.text:
                raise ValueError('Replacement has no changes')
            signatures = {str(run._r.rPr.xml) if run._r.rPr is not None else '' for run in para.runs if run.text}
            if len(signatures) > 1:
                raise ValueError('Mixed run formatting requires manual polishing')
            # Language changes cannot silently change quantitative claims or citations.
            if re.findall(r'\d+(?:\.\d+)?', text) != re.findall(r'\d+(?:\.\d+)?', para.text):
                raise ValueError('Polishing changed numbers; handle factual changes separately')
            if [m[0] for m in CITATION_RE.finditer(text)] != [m[0] for m in CITATION_RE.finditer(para.text)]:
                raise ValueError('Polishing changed citation markers')
            prepared.append((kind, para, text))
            changes.append({'op': kind, 'paragraph_index': index, 'before': para.text, 'after': text})
        elif kind == 'format':
            # Validate on a detached copy; no source paragraph is changed before preflight completes.
            import copy
            from docx.text.paragraph import Paragraph
            _format(Paragraph(copy.deepcopy(para._p), doc), op.get('formatting'))
            prepared.append((kind, para, op['formatting']))
            changes.append({'op': kind, 'paragraph_index': index, 'formatting': op['formatting']})
        else:
            text = _text(op.get('text'), 'citation text')
            ref_id = _text(op.get('ref_id'), 'ref_id')
            ref = references.get(ref_id)
            if not isinstance(ref, dict):
                raise ValueError(f'Unknown reference {ref_id}')
            entry = _text(ref.get('entry'), 'full reference entry')
            quote = _text(ref.get('quote'), 'source quote')
            _text(ref.get('support_note'), 'host semantic support note')
            evidence_path = (base_dir / _text(ref.get('source'), 'source')).resolve()
            if evidence_path not in evidence_cache:
                evidence_cache[evidence_path] = read_evidence(str(evidence_path))
            evidence = evidence_cache[evidence_path]
            if ref.get('source_sha256') != evidence['source_sha256']:
                raise ValueError(f'Evidence source changed for {ref_id}')
            page = ref.get('page')
            if type(page) is not int or not 1 <= page <= len(evidence['pages']):
                raise ValueError(f'Unknown evidence page for {ref_id}')
            normalize = lambda value: re.sub(r'\s+', ' ', value).strip()
            if normalize(quote) not in normalize(evidence['pages'][page - 1]['text']):
                raise ValueError(f'Quote does not occur on the supplied page for {ref_id}')
            positions.append((para, CitationEdit('', text, entry, '', '', entry,
                reason=ref['support_note'], existing_number=ref.get('existing_number'))))
            changes.append({'op': kind, 'paragraph_index': index, 'text': text,
                'ref_id': ref_id, 'source': str(evidence_path), 'source_sha256': evidence['source_sha256'],
                'page': page, 'quote': quote, 'support_note': ref['support_note']})
    for kind, para, value in prepared:
        if kind == 'replace':
            if para.runs:
                para.runs[0].text = value
                for run in para.runs[1:]:
                    run.text = ''
            else:
                para.add_run(value)
        else:
            _format(para, value)
    ref_format = plan.get('reference_format')
    bibliography = []
    if positions or ref_format is not None:
        if not positions and ref_start == len(paras):
            raise ValueError('No bibliography found to format')
        style = get_builtin_style('gbt7714-numeric')
        if ref_format is not None:
            # Same small formatting vocabulary for bibliography and body paragraphs.
            probe = Document().add_paragraph('probe')
            _format(probe, ref_format)
            for language in ('cn', 'en'):
                style.font_rule[language]['name'] = ref_format.get('font_' + language, style.font_rule[language]['name'])
                style.font_rule[language]['size'] = ref_format.get('size_pt', style.font_rule[language]['size'])
            for name in ('line_spacing', 'space_before', 'space_after'):
                key = name if name == 'line_spacing' else name + '_pt'
                if key in ref_format:
                    style.spacing_rule[name] = ref_format[key]
            if 'first_line_indent_pt' in ref_format:
                value = ref_format['first_line_indent_pt']
                style.indent_rule = {'type': 'hanging' if value < 0 else 'first_line', 'value': f'{abs(value)}pt'}
            if 'left_indent_pt' in ref_format:
                raise ValueError('Reference indent is controlled by first_line_indent_pt; omit left_indent_pt')
        bibliography, _ = DocxEditor._apply_numeric_citations(doc, positions, style)
        if ref_format is not None:
            changes.append({'op': 'reference_format', 'formatting': ref_format})
    if str(doc.element.xml) == original_xml:
        raise ValueError('The requested operations made no changes')
    requested = len(operations) + int(ref_format is not None)
    result = {'schema_version': 1, 'status': 'validated' if validate_only else 'ready_for_review',
              'source_sha256': source_hash, 'requested': requested, 'applied': 0 if validate_only else requested,
              'checks': {'anchors_and_source': True, 'source_quotes': True,
                         'semantic_review': 'Host assessment recorded; user review still required',
                         'user_accepted': False}, 'changes': changes, 'bibliography': bibliography, 'artifacts': {}}
    # Source stability checked a second time to catch changes during evidence processing.
    if hashlib.sha256(source_path.read_bytes()).hexdigest() != source_hash:
        raise ValueError('Source changed during processing')
    if validate_only:
        return result
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = Path(tempfile.mkdtemp(prefix='.paper-attempt-', dir=target.parent))
    try:
        DocxEditor.save_new(doc, temporary / 'result.docx')
        (temporary / 'input.docx').write_bytes(raw)
        (temporary / 'edits.json').write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding='utf-8')
        result['artifacts'] = {key: str(target / filename) for key, filename in {
            'document': 'result.docx', 'input': 'input.docx', 'edits': 'edits.json', 'report': 'report.json'}.items()}
        (temporary / 'report.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
        if target.exists():
            raise ValueError('Attempt directory appeared during processing')
        os.rename(temporary, target)
    finally:
        if temporary.exists() and temporary.resolve().parent == target.parent:
            shutil.rmtree(temporary)
    return result
