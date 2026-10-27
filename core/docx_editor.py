"""Document Editor — Phase 3-4.

Insert citations into .docx, renumber references in order of first appearance,
and export enhanced paper + citation report.
"""

import copy
import os
import re
from io import BytesIO
from docx.text.paragraph import Paragraph
from .docx_parser import (CITATION_RE, citation_numbers, is_reference_heading,
                          is_post_reference_heading, plain_paragraph)
from dataclasses import dataclass, field
from docx.oxml.ns import qn

from .style_config import ReferenceStyleConfig, get_builtin_style
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Emu


@dataclass
class CitationEdit:
    """Describes one citation to insert."""
    section_heading: str  # target section heading
    citation_text: str    # the generated citation sentence
    ref_title: str        # reference title
    ref_authors: str      # reference authors (for bibliography)
    ref_year: str         # reference year
    ref_full: str         # full reference entry for bibliography
    confidence: float = 1.0
    reason: str = ""
    existing_number: int | None = None


@dataclass
class EditResult:
    output_path: str
    edits_applied: int
    new_references: list[str]  # ordered by first appearance
    report_rows: list[dict]


class DocxEditor:
    """Edit .docx: insert citations, renumber, export."""

    def __init__(self):
        pass

    def apply_edits(
        self,
        docx_path: str,
        edits: list[CitationEdit],
        output_path: str,
        style_config: ReferenceStyleConfig = None,
    ) -> EditResult:
        """Apply citation edits to a .docx and export enhanced version.

        Args:
            docx_path: Path to original paper .docx.
            edits: List of CitationEdit to apply (one per citation insertion).
            output_path: Where to save enhanced_paper.docx.
            style_config: ReferenceStyleConfig (defaults to gbt7714-numeric).

        Returns:
            EditResult with output path, edit count, new reference list, report rows.
        """
        if style_config is None:
            style_config = get_builtin_style("gbt7714-numeric")

        self.check_output(docx_path, output_path)
        if not edits:
            raise ValueError('No edits requested')
        doc = Document(docx_path)
        positions = []
        for edit in edits:
            matches = [i for i, p in enumerate(doc.paragraphs)
                       if ' '.join(p.text.split()) == ' '.join(edit.section_heading.split())]
            if len(matches) != 1:
                raise ValueError(f'Heading must match exactly once: {edit.section_heading}')
            end = self._find_section_end(doc, matches[0])
            positions.append((doc.paragraphs[end - 1], edit))
        refs, reports = self._apply_numeric_citations(doc, positions, style_config)
        self.save_new(doc, output_path)
        return EditResult(str(Path(output_path).resolve()), len(edits), refs, reports)

    @staticmethod
    def check_output(source, output):
        if Path(source).resolve() == Path(output).resolve() or Path(output).exists():
            raise ValueError('Output must be a new path, never an input or previous result')

    @staticmethod
    def save_new(doc, output):
        # Serialize and reopen BEFORE publishing; exclusive creation prevents overwrites.
        buffer = BytesIO()
        doc.save(buffer)
        raw = buffer.getvalue()
        Document(BytesIO(raw))
        path = Path(output)
        path.parent.mkdir(parents=True, exist_ok=True)
        try:
            with path.open('xb') as stream:
                stream.write(raw)
        except FileExistsError:
            raise ValueError('Output already exists; choose a new attempt')

    @staticmethod
    def _replace_span(para, start, end, replacement):
        if not plain_paragraph(para):
            raise ValueError('Cannot safely rewrite a citation inside a complex paragraph')
        offset = 0
        first = True
        for run in para.runs:
            text = run.text
            stop = offset + len(text)
            if offset < end and stop > start:
                left, right = max(0, start - offset), min(len(text), end - offset)
                run.text = text[:left] + (replacement if first else '') + text[right:]
                first = False
            offset = stop

    @staticmethod
    def _apply_ref_format(para, style):
        for run in para.runs:
            font = style.font_rule['en' if DocxEditor._is_english_ref(para.text) else 'cn']
            run.font.name = font['name']
            run.font.size = Pt(font['size'])
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), style.font_rule['cn']['name'])
        pf = para.paragraph_format
        spacing = style.spacing_rule
        pf.line_spacing = spacing.get('line_spacing', 1.5)
        pf.space_before = Pt(spacing.get('space_before', 0))
        pf.space_after = Pt(spacing.get('space_after', 0))
        indent = style.indent_rule
        value = str(indent.get('value', '2char'))
        match = re.fullmatch(r'(\d+(?:\.\d+)?)(char|pt)', value)
        if not match:
            raise ValueError('Indent must use char or pt units')
        points = float(match[1]) * (style.font_rule['cn']['size'] if match[2] == 'char' else 1)
        pf.left_indent = Pt(points if indent['type'] == 'hanging' else 0)
        pf.first_line_indent = Pt(-points if indent['type'] == 'hanging' else points if indent['type'] == 'first_line' else 0)

    @staticmethod
    def _apply_numeric_citations(doc, positions, style):
        if (style.numbering_rule != 'bracket_numeric' or style.ordering_rule != 'first_appearance'
                or style.language_split_rule != 'mixed'):
            raise ValueError('Automatic editing currently supports numeric first-appearance order only')
        # Fields and revision containers may hide citations from python-docx text.
        forbidden = ['fldChar', 'fldSimple', 'instrText', 'ins', 'del', 'footnoteReference', 'endnoteReference', 'txbxContent', 'altChunk']
        if any(doc.element.body.xpath('.//w:' + name) for name in forbidden):
            raise ValueError('Fields, notes or tracked changes require manual citation handling')
        for part in doc.part.package.parts:
            if any(word in str(part.partname) for word in ('header', 'footer', 'footnotes', 'endnotes')):
                if re.search(rb'\[\d', part.blob):
                    raise ValueError('Citations outside the document body require manual handling')
        starts = [i for i, p in enumerate(doc.paragraphs) if is_reference_heading(p.text)]
        if len(starts) > 1:
            raise ValueError('Multiple bibliographies require an explicit manual boundary')
        old, end = [], None
        if starts:
            start = starts[0]
            old, end, _ = DocxEditor._parse_old_refs(doc, start)
            ref_heading = doc.paragraphs[start]
            boundary = doc.paragraphs[end]._p if end < len(doc.paragraphs) else None
            node = ref_heading._p.getnext()
            while node is not None and node is not boundary:
                if node.tag not in {qn('w:p'), qn('w:sectPr')}:
                    raise ValueError('Non-paragraph bibliography content requires manual handling')
                node = node.getnext()
        else:
            ref_heading = None
        records = {r['number']: r for r in old}
        old_elements = {r['paragraph']._p for r in old}
        all_paras = [Paragraph(el, doc) for el in doc.element.body.iter(qn('w:p')) if el not in old_elements]
        # Check existing references before adding provisional numbers.
        for para in all_paras:
            if re.search(r'[［【]\s*\d', para.text):
                raise ValueError('Unsupported citation brackets require manual review')
            spans = [(m.start(), m.end()) for m in CITATION_RE.finditer(para.text)]
            offset = 0
            for run in para.runs:
                if run.font.superscript:
                    for digit in re.finditer(r'\d', run.text):
                        if not any(a <= offset + digit.start() < b for a, b in spans):
                            raise ValueError('Unbracketed superscript numbers require manual review')
                offset += len(run.text)
            for candidate in re.finditer(r'\[\s*\d[^\]]*\]', para.text):
                if not CITATION_RE.fullmatch(candidate[0]):
                    raise ValueError('Unsupported numeric bracket syntax requires manual review')
            for match in CITATION_RE.finditer(para.text):
                if not plain_paragraph(para):
                    raise ValueError('Complex citation paragraph requires manual handling')
                if any(n not in records for n in citation_numbers(match[1])):
                    raise ValueError('An existing citation has no corresponding bibliography entry')
        reports, last_inserted = [], {}
        next_number = max(records, default=0) + 1
        body_indices = {p._p: i for i, p in enumerate(doc.paragraphs)}
        for anchor, edit in positions:
            if anchor._p not in body_indices or (starts and body_indices[anchor._p] >= starts[0]):
                raise ValueError('Insertions must target body paragraphs before the bibliography')
            if is_post_reference_heading(anchor.text) or not plain_paragraph(anchor):
                raise ValueError('Unsupported insertion anchor')
            if not edit.citation_text.strip() or CITATION_RE.search(edit.citation_text) or '生成失败' in edit.citation_text:
                raise ValueError('Supply successful citation text without numeric markers')
            raw = re.sub(r'^\[\d+\]\s*', '', edit.ref_full).strip()
            if not raw:
                raise ValueError('Full reference entry is required')
            norm = lambda text: ' '.join(text.split()).casefold()
            matches = [n for n, r in records.items() if norm(r['full_text']) == norm(raw)]
            if edit.existing_number is not None:
                number = edit.existing_number
                if number not in records or number not in matches:
                    raise ValueError('existing_number must identify the same full reference')
            elif len(matches) > 1:
                raise ValueError('Duplicate bibliography entries require existing_number')
            elif matches:
                number = matches[0]
            else:
                number = next_number
                next_number += 1
                records[number] = {'full_text': raw, 'paragraph': None}
            para = doc.add_paragraph()
            if anchor._p.pPr is not None:
                para._p.insert(0, copy.deepcopy(anchor._p.pPr))
            run = para.add_run(f'{edit.citation_text}[{number}]')
            if anchor.runs and anchor.runs[-1]._r.rPr is not None:
                run._r.insert(0, copy.deepcopy(anchor.runs[-1]._r.rPr))
            target = last_inserted.get(anchor._p, anchor._p)
            target.addnext(para._p)
            last_inserted[anchor._p] = para._p
            reports.append({'插入位置': anchor.text, '引用内容': edit.citation_text,
                            '来源文献': raw, '页码': '', '置信度': edit.confidence, '插入原因': edit.reason})
        # Number from FINAL body order, including table paragraphs and appendices.
        ordered = []
        paragraphs = [Paragraph(el, doc) for el in doc.element.body.iter(qn('w:p')) if el not in old_elements]
        for para in paragraphs:
            for match in CITATION_RE.finditer(para.text):
                for number in citation_numbers(match[1]):
                    if number not in ordered:
                        ordered.append(number)
        ordered.extend(n for n in records if n not in ordered)  # preserve uncited old entries
        mapping = {old_number: i for i, old_number in enumerate(ordered, 1)}
        for para in paragraphs:
            for match in reversed(list(CITATION_RE.finditer(para.text))):
                text = '[' + ','.join(str(mapping[n]) for n in citation_numbers(match[1])) + ']'
                if text != match[0]:
                    DocxEditor._replace_span(para, match.start(), match.end(), text)
        if ref_heading is None:
            ref_heading = doc.add_paragraph('参考文献')
            tail = next((p for p in doc.paragraphs if is_post_reference_heading(p.text)), None)
            if tail is not None:
                tail._p.addprevious(ref_heading._p)
        base = ref_heading._p
        full_refs = []
        for i, number in enumerate(ordered, 1):
            record = records[number]
            para = record.get('paragraph')
            if para is None:
                para = doc.add_paragraph(f"[{i}] {record['full_text']}")
            else:
                marker = re.match(r'^\s*\[\d+\]', para.text)
                DocxEditor._replace_span(para, marker.start(), marker.end(), f'[{i}]')
            DocxEditor._apply_ref_format(para, style)
            base.addnext(para._p)
            base = para._p
            full_refs.append(record['full_text'])
        return full_refs, reports

    @staticmethod
    def _map_headings_to_paragraphs(doc: Document) -> dict[str, int]:
        """Map section heading text → paragraph index. Normalizes whitespace for matching."""
        mapping = {}
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name if para.style else ""
            if is_reference_heading(text) or is_post_reference_heading(text) or style.startswith("Heading") or re.match(r"^\d+\.", text) or re.match(r"^第[一二三四五六七八九十\d]+章", text):
                # Store both normalized and original to handle double-space issues
                normalized = re.sub(r"\s+", " ", text)
                mapping[text] = i
                mapping[normalized] = i
        return mapping

    @staticmethod
    def _find_section_end(doc: Document, start_idx: int) -> int:
        """Find the paragraph index after the current section's last paragraph."""
        for i in range(start_idx + 1, len(doc.paragraphs)):
            text = doc.paragraphs[i].text.strip()
            style = doc.paragraphs[i].style.name if doc.paragraphs[i].style else ""
            if is_reference_heading(text) or is_post_reference_heading(text) or style.startswith("Heading") or re.match(r"^\d+\.", text) or re.match(r"^第[一二三四五六七八九十\d]+章", text):
                return i
        return len(doc.paragraphs)

    @staticmethod
    def _find_or_add_ref(
        refs: list[tuple[str, str, str, str]],
        title: str,
        authors: str,
        year: str,
        full: str,
    ) -> int:
        """Find existing ref or add new one. Returns 1-based index by first appearance."""
        for i, (t, _, _, _) in enumerate(refs):
            if t == title:
                return i + 1
        refs.append((title, authors, year, full))
        return len(refs)

    @staticmethod
    def _build_citation_text(citation_text: str, ref_index: int) -> str:
        """Build citation sentence with marker."""
        # If citation_text already has a marker like [1], don't double-add
        if re.search(r"\[\d+\]", citation_text):
            return citation_text
        return f"{citation_text}[{ref_index}]"

    @staticmethod
    def _insert_citation_paragraph(
        doc: Document,
        insert_pos: int,
        citation_text: str,
        confidence: float,
        style_config: ReferenceStyleConfig,
    ):
        """Insert a citation paragraph at the given position."""
        ref_element = doc.paragraphs[insert_pos - 1]._element if insert_pos > 0 else doc.element.body

        para = doc.add_paragraph()
        if insert_pos < len(doc.paragraphs):
            target = doc.paragraphs[insert_pos]._element
            target.addprevious(para._element)
        else:
            ref_element.addnext(para._element)

        cn_font = style_config.font_rule["cn"]
        run = para.add_run(citation_text)
        run.font.size = Pt(cn_font["size"])
        run.font.name = cn_font["name"]
        run._element.rPr.rFonts.set(qn("w:eastAsia"), cn_font["name"])

        if confidence < 0.85:
            note = para.add_run(f" [AI推荐，置信度{confidence:.0%}，请核实]")
            note.font.size = Pt(9)
            note.font.color.rgb = RGBColor(150, 150, 150)

    # Paragraphs that signal the end of the references section
    _POST_REF_MARKERS = [
        re.compile(r"^致谢"),
        re.compile(r"^Acknowledgements?"),
        re.compile(r"^附录"),
        re.compile(r"^Appendix"),
        re.compile(r"^声明"),
        re.compile(r"^Declaration"),
        re.compile(r"^作者简介"),
        re.compile(r"^研究成果"),
        re.compile(r"^在读期间"),
        re.compile(r"^攻读.*期间"),
        re.compile(r"^已发表"),
        re.compile(r"^发表论文"),
    ]

    @staticmethod
    def _parse_old_refs(doc: Document, ref_start: int) -> tuple[list[dict], int, dict | None]:
        """Parse old reference paragraphs from the docx.

        Returns (old_refs, ref_end_idx, font_template).
        old_refs: [{title, full_text}]
        font_template: {name, size} from first ref paragraph, or None
        """
        old_refs = []
        ref_end = ref_start + 1
        font_template = None

        for i in range(ref_start + 1, len(doc.paragraphs)):
            text = doc.paragraphs[i].text.strip()
            if not text:
                if not plain_paragraph(doc.paragraphs[i]):
                    raise ValueError('Non-text bibliography content requires manual handling')
                ref_end = i + 1
                continue
            style = doc.paragraphs[i].style.name if doc.paragraphs[i].style else ""

            # Stop at post-reference markers
            if is_post_reference_heading(text):
                break
            if style.startswith("Heading") and len(text) < 30:
                break

            number_match = re.match(r'^\[(\d+)\]\s*(\S.*)$', text)
            if not number_match or int(number_match[1]) < 1:
                raise ValueError(f'Ambiguous bibliography boundary or unsupported entry: paragraph {i}')
            if any(r['number'] == int(number_match[1]) for r in old_refs):
                raise ValueError('Duplicate bibliography number')
            ref_end = i + 1
            para = doc.paragraphs[i]
            if not plain_paragraph(para):
                raise ValueError('Complex bibliography entries require manual handling')

            # Capture formatting from first reference paragraph
            if font_template is None and para.runs:
                r = para.runs[0]
                font_template = {
                    "name": r.font.name or "宋体",
                    "size": r.font.size or Pt(12),
                }

            # Extract title from this ref entry; strip old [N] prefix
            raw = re.sub(r"^\[\d+\]\s*", "", text).strip()
            title = DocxEditor._extract_title_from_ref(raw)
            old_refs.append({
                "number": int(number_match[1]),
                "paragraph": para,
                "title": title or raw[:80],
                "full_text": raw,  # without old [N] prefix
            })

        return old_refs, ref_end, font_template

    @staticmethod
    def _extract_title_from_ref(ref_text: str) -> str:
        """Extract paper title from a GB/T 7714 reference string."""
        ref_norm = ref_text.replace("．", ".").replace("、", ",")
        # Chinese: author. TITLE[J/M/D/C/P]...
        m = re.match(r"^[^\]\.]+[\.．](.+?)\[[JMCDP]\]", ref_norm)
        if m:
            return m.group(1).strip()
        # English: Author, N. TITLE[M]...
        m = re.match(r"^[A-Z][^\.]+\.[\s．]*(.+?)\[[JMCDP]\]", ref_norm, re.DOTALL)
        if m:
            return m.group(1).strip()
        # Fallback: return first 80 chars
        return ref_text[:80]

    @staticmethod
    def _is_english_ref(text: str) -> bool:
        """True if reference text contains no CJK characters."""
        return not bool(re.search(r"[一-鿿㐀-䶿]", text))

    @staticmethod
    def _extract_author_sort_key(ref_text: str) -> str:
        """Extract the first author's last name for alpha sorting."""
        ref_clean = re.sub(r"^\[\d+\]\s*", "", ref_text).strip()
        # English: "Smith, J. ..." or "Smith J. ..."
        m = re.match(r"^([A-Z][a-z]+)", ref_clean)
        if m:
            return m.group(1).lower()
        # Chinese: use first few characters as proxy
        m = re.match(r"^([一-鿿]{1,4})", ref_clean)
        if m:
            return m.group(1)
        return ref_clean[:10].lower()

    @staticmethod
    def _normalize(s: str) -> str:
        return re.sub(r"[^\w一-鿿]", "", s.lower())

    @staticmethod
    def _fuzzy_match(a: str, b: str, threshold: float = 0.5) -> bool:
        if not a or not b:
            return False
        if a == b:
            return True
        if a in b or b in a:
            return True
        chars_a, chars_b = set(a), set(b)
        if not chars_a or not chars_b:
            return False
        overlap = len(chars_a & chars_b) / min(len(chars_a), len(chars_b))
        return overlap > threshold


def export_citation_report(rows: list[dict], output_path: str) -> None:
    """Export citation report to .xlsx for manual review."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "引用报告"

    headers = ["插入位置", "引用内容", "来源文献", "页码", "置信度", "插入原因"]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    wrap = Alignment(wrap_text=True, vertical="top")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = wrap

    for i, row in enumerate(rows, 1):
        for col, key in enumerate(headers, 1):
            cell = ws.cell(row=i + 1, column=col, value=row.get(key, ""))
            cell.alignment = wrap

    for col_cells in ws.columns:
        max_len = max(min(len(str(cell.value or "")), 60) for cell in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = max_len + 2

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    # Preserve existing artifacts, including under concurrent invocations.
    buffer = BytesIO()
    wb.save(buffer)
    with open(output_path, 'xb') as stream:
        stream.write(buffer.getvalue())
