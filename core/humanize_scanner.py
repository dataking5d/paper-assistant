"""Read-only AI-writing-trace scanner.

Design rules:
- report candidates, never replacement text;
- preserve numbers, citations, hedges, negations and attributions;
- technical terms are not errors by themselves;
- headings, captions, references, tables, headers/footers are skipped.
"""
from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any

import docx
import yaml

from .docx_parser import text_hash

RULES_PATH = (
    Path(__file__).resolve().parents[1]
    / "capabilities"
    / "humanizing"
    / "references"
    / "scanner-rules.yaml"
)


def sha256_file(path: str | Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def load_rules(path: str | Path | None = None) -> dict[str, Any]:
    p = Path(path) if path else RULES_PATH
    return yaml.safe_load(p.read_text(encoding="utf-8"))


def detect_language(text: str) -> str:
    zh = len(re.findall(r"[\u4e00-\u9fff]", text))
    en = len(re.findall(r"[A-Za-z]", text))
    return "zh" if zh >= en else "en"


def _spans(text: str, pattern: str) -> list[tuple[int, int, str]]:
    try:
        return [(m.start(), m.end(), m.group(0)) for m in re.finditer(pattern, text)]
    except re.error:
        return []


def _overlaps(spans: list[tuple[int, int, str]], start: int, end: int) -> bool:
    return any(s < end and start < e for s, e, _ in spans)


def preservation_snapshot(text: str, protection: dict[str, str]) -> dict[str, list[str]]:
    return {
        "numbers": re.findall(protection.get("numbers", r"\d+"), text),
        "citations": re.findall(protection.get("citations", r"\[\d+\]"), text),
        "negations": re.findall(protection.get("negations", r"不"), text),
        "hedges": re.findall(protection.get("hedges", r"可能"), text),
        "attributions": re.findall(protection.get("attributions", r"根据"), text),
    }


def _finding(
    rule_id: str,
    category: str,
    severity: str,
    confidence: float,
    span: tuple[int, int],
    matched: str,
    reason: str,
    action: str,
    protected: bool,
) -> dict[str, Any]:
    return {
        "rule_id": rule_id,
        "category": category,
        "severity": severity,
        "confidence": round(confidence, 2),
        "span": [int(span[0]), int(span[1])],
        "matched_text": matched,
        "reason": reason,
        "suggested_action": action,
        "auto_fixable": False,
        "requires_agent": True,
        "protected": bool(protected),
    }


def scan_text(text: str, language: str | None = None, rules: dict[str, Any] | None = None) -> dict[str, Any]:
    rules = rules or load_rules()
    lang = language or detect_language(text)
    pack = rules.get(lang) or rules.get("zh") or {}
    protection = rules.get("protection") or {}
    snap = preservation_snapshot(text, protection)
    protected_spans = {
        "numbers": _spans(text, protection.get("numbers", r"\d+")),
        "citations": _spans(text, protection.get("citations", r"\[\d+\]")),
        "hedges": _spans(text, protection.get("hedges", r"可能")),
        "attributions": _spans(text, protection.get("attributions", r"根据")),
    }
    all_protected = [s for group in protected_spans.values() for s in group]
    findings: list[dict[str, Any]] = []

    # Weak word lists: one hit is low; three or more in one paragraph is medium.
    whitelist = pack.get("technical_whitelist", []) or []
    word_hits = []
    for word in pack.get("words", []) or []:
        if any(w in text and word in w for w in whitelist):
            continue
        for start, end, matched in _spans(text, re.escape(word)):
            word_hits.append((start, end, matched))
    for start, end, matched in word_hits:
        severity = "medium" if len(word_hits) >= 3 else "low"
        findings.append(_finding(
            f"{lang}.pattern.12.ai_word", "word_choice", severity, 0.55, (start, end), matched,
            "AI 高频词候选；只有空泛、重复或用词不准时才改", "keep",
            _overlaps(all_protected, start, end),
        ))

    for phrase in pack.get("inflated", []) or []:
        for start, end, matched in _spans(text, re.escape(phrase)):
            findings.append(_finding(
                f"{lang}.pattern.13.inflated", "inflation", "high", 0.75, (start, end), matched,
                "意义拔高候选；删掉后若不影响信息就删，不能补造新论点", "delete_filler",
                _overlaps(all_protected, start, end),
            ))

    for phrase in pack.get("vague_authority", []) or []:
        for start, end, matched in _spans(text, re.escape(phrase)):
            nearby_cite = any(abs(s - start) <= 40 for s, _, _ in protected_spans["citations"])
            severity = "medium" if nearby_cite else "high"
            reason = "已有引用/来源，核对是否真的支持该句" if nearby_cite else "模糊归因候选；需要具体来源，不能补造"
            findings.append(_finding(
                f"{lang}.pattern.17.vague_authority", "attribution", severity, 0.7, (start, end), matched,
                reason, "verify_citation", _overlaps(all_protected, start, end),
            ))

    for phrase in pack.get("closing", []) or []:
        for start, end, matched in _spans(text, re.escape(phrase)):
            findings.append(_finding(
                f"{lang}.pattern.31.closing_formula", "closure", "medium", 0.7, (start, end), matched,
                "套话收尾候选；有实际内容的总结可保留", "delete_filler",
                _overlaps(all_protected, start, end),
            ))

    for rule in pack.get("regex", []) or []:
        for start, end, matched in _spans(text, rule["pattern"]):
            protected = _overlaps(all_protected, start, end)
            severity = "error" if (protected and not rule.get("no_escalate")) else rule.get("severity", "low")
            reason = rule.get("reason", "")
            if protected:
                reason += "；命中数字/引用/限定/归因，必须人工复核"
            findings.append(_finding(
                rule["id"], rule.get("category", "pattern"), severity,
                rule.get("confidence", 0.6), (start, end), matched,
                reason, rule.get("action", "rewrite_clause"), protected,
            ))

    # stable order; de-duplicate identical rule+span
    unique = {}
    for f in findings:
        key = (f["rule_id"], tuple(f["span"]))
        unique[key] = f
    findings = sorted(unique.values(), key=lambda f: (f["span"][0], f["rule_id"]))
    return {"language": lang, "findings": findings, "preservation_snapshot": snap}


def _is_skippable(paragraph: docx.text.paragraph.Paragraph) -> tuple[bool, str]:
    style = (paragraph.style.name if paragraph.style else "") or ""
    text = (paragraph.text or "").strip()
    if not text:
        return True, "empty"
    if style.lower().startswith("heading"):
        return True, "heading"
    if style.lower().startswith("caption"):
        return True, "caption"
    if text.startswith("表 ") or text.startswith("图 ") or text.startswith("数据来源"):
        return True, "caption_or_source"
    if text.startswith("参考文献") or text.lower().startswith("references"):
        return True, "references_heading"
    if re.fullmatch(r"[\d\W_]+", text):
        return True, "numbers_or_punctuation"
    try:
        if paragraph._p.xpath(".//w:drawing"):
            return True, "figure"
    except Exception:
        pass
    return False, ""


def scan_docx(path: str | Path, language: str | None = None, rules: dict[str, Any] | None = None) -> dict[str, Any]:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"docx not found: {p}")
    rules = rules or load_rules()
    doc = docx.Document(str(p))
    paragraphs = []
    section = ""
    in_references = False
    for idx, para in enumerate(doc.paragraphs):
        text = (para.text or "").strip()
        style = (para.style.name if para.style else "") or ""
        if style.lower().startswith("heading"):
            section = text
            continue
        if text.startswith("参考文献") or text.lower().startswith("references"):
            in_references = True
            continue
        if in_references:
            continue
        skip, _reason = _is_skippable(para)
        if skip:
            continue
        result = scan_text(text, language=language, rules=rules)
        if result["findings"]:
            paragraphs.append({
                "paragraph_index": idx,
                "section": section,
                "text_sha256": text_hash(para.text),
                "findings": result["findings"],
                "preservation_snapshot": result["preservation_snapshot"],
            })
    by_severity: dict[str, int] = {}
    by_category: dict[str, int] = {}
    for para in paragraphs:
        for f in para["findings"]:
            by_severity[f["severity"]] = by_severity.get(f["severity"], 0) + 1
            by_category[f["category"]] = by_category.get(f["category"], 0) + 1
    return {
        "schema_version": 1,
        "language": language or "auto",
        "ruleset_version": rules.get("version", "unknown"),
        "document_sha256": sha256_file(p),
        "summary": {
            "paragraphs_scanned": len([p for p in doc.paragraphs if (p.text or "").strip()]),
            "paragraphs_with_findings": len(paragraphs),
            "findings": sum(len(p["findings"]) for p in paragraphs),
            "by_severity": by_severity,
            "by_category": by_category,
        },
        "paragraphs": paragraphs,
        "skipped": ["heading", "caption", "figure", "table", "header", "footer", "footnote", "references_section"],
    }
