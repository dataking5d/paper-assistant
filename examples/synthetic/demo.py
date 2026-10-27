"""Run the offline CLI on synthetic DOCX content, without a model or network."""

import argparse
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
ORIGINAL = "这不是一个工具，而是一个平台。研究表明，该方案有效。"
REPLACEMENT = "这个演示用于展示文档解析、修改清单和版本保存。"
TAIL = "这段合成文字用于检查未修改的内容是否保留。"


def cli(*args):
    result = subprocess.run(
        [sys.executable, "-B", "-X", "utf8", str(ROOT / "main.py"), *args],
        cwd=ROOT, capture_output=True, text=True, encoding="utf-8",
    )
    if result.returncode:
        raise RuntimeError(result.stderr or result.stdout)
    return json.loads(result.stdout)


def save_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def require(condition, message):
    if not condition:
        raise RuntimeError(message)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, help="New output directory; must not exist")
    args = parser.parse_args()
    if args.output_dir:
        output = args.output_dir.resolve()
        output.mkdir(parents=True, exist_ok=False)
    else:
        output = Path(tempfile.mkdtemp(prefix="paper_assistant_demo_"))

    source = output / "synthetic.docx"
    doc = Document()
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(ORIGINAL)
    doc.add_paragraph(TAIL)
    doc.save(source)
    original_hash = hashlib.sha256(source.read_bytes()).hexdigest()

    snapshot = cli("parse", str(source), "--json")
    save_json(output / "parse.json", snapshot)
    target = next(item for item in snapshot["paragraphs"] if item["text"] == ORIGINAL)
    scan = cli("humanize-scan", str(source), "--language", "zh", "--json")
    save_json(output / "humanize_scan.json", scan)
    require(scan["document_sha256"] == original_hash, "Scan fingerprint differs from input")

    plan = {
        "schema_version": 1,
        "source_sha256": snapshot["source_sha256"],
        "operations": [{
            "op": "replace", "paragraph_index": target["index"],
            "text_sha256": target["text_sha256"], "text": REPLACEMENT,
        }],
        "references": {},
    }
    plan_path = output / "edits.json"
    save_json(plan_path, plan)
    attempt = output / "attempt"
    edit_args = ("edit", "--paper", str(source), "--edits", str(plan_path), "--attempt", str(attempt))
    validation = cli(*edit_args, "--validate-only")
    save_json(output / "validation.json", validation)
    require(validation["status"] == "validated", "Preflight did not validate")
    require(not attempt.exists(), "Preflight created an attempt directory")

    result = cli(*edit_args)
    require(result["status"] == "ready_for_review", "Edit did not produce a reviewable result")
    edited = Document(result["artifacts"]["document"])
    require(edited.paragraphs[target["index"]].text == REPLACEMENT, "Replacement differs")
    require(edited.paragraphs[-1].text == TAIL, "Untouched paragraph differs")
    require(hashlib.sha256(source.read_bytes()).hexdigest() == original_hash, "Input was modified")
    require(all((attempt / name).is_file() for name in ("input.docx", "edits.json", "result.docx", "report.json")),
            "Attempt artifacts are incomplete")

    summary = {
        "status": "passed", "synthetic_data_only": True, "input_unchanged": True,
        "preflight_wrote_no_attempt": True, "replacement_verified": True,
        "untouched_paragraph_preserved": True, "output_dir": str(output),
        "review_document": result["artifacts"]["document"],
        "scope": "CLI/file checks only; visual and semantic review remain manual",
    }
    save_json(output / "demo_summary.json", summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
